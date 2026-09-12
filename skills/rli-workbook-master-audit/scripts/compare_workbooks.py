#!/usr/bin/env python3
"""Compare values/formulas in common worksheets without modifying either file."""

from __future__ import annotations

import argparse
import json
from itertools import zip_longest

import openpyxl
from openpyxl.utils import get_column_letter


def clean(value):
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if hasattr(value, "text") or hasattr(value, "ref"):
        return {
            "formula_type": type(value).__name__,
            "text": getattr(value, "text", None),
            "ref": getattr(value, "ref", None),
        }
    return str(value)


def compare_sheet(left, right, example_limit: int) -> dict:
    different = left_only = right_only = 0
    examples = []
    left_rows = left.iter_rows(values_only=True)
    right_rows = right.iter_rows(values_only=True)
    for row_num, (lrow, rrow) in enumerate(zip_longest(left_rows, right_rows, fillvalue=()), start=1):
        for col_num, (lval, rval) in enumerate(zip_longest(lrow, rrow, fillvalue=None), start=1):
            left_value = clean(lval)
            right_value = clean(rval)
            if left_value == right_value:
                continue
            different += 1
            if lval is not None and rval is None:
                left_only += 1
            elif rval is not None and lval is None:
                right_only += 1
            if len(examples) < example_limit:
                examples.append({
                    "cell": f"{get_column_letter(col_num)}{row_num}",
                    "left": left_value,
                    "right": right_value,
                })
    return {
        "left_dimension": left.calculate_dimension(),
        "right_dimension": right.calculate_dimension(),
        "different_cells": different,
        "left_only_nonblank": left_only,
        "right_only_nonblank": right_only,
        "examples": examples,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("left")
    parser.add_argument("right")
    parser.add_argument("--examples", type=int, default=5)
    parser.add_argument("--output")
    args = parser.parse_args()
    left = openpyxl.load_workbook(args.left, read_only=True, data_only=False, keep_links=False)
    right = openpyxl.load_workbook(args.right, read_only=True, data_only=False, keep_links=False)
    result = {
        "left": args.left,
        "right": args.right,
        "left_only_sheets": [name for name in left.sheetnames if name not in right.sheetnames],
        "right_only_sheets": [name for name in right.sheetnames if name not in left.sheetnames],
        "common_sheets": {},
    }
    for name in left.sheetnames:
        if name in right.sheetnames:
            result["common_sheets"][name] = compare_sheet(left[name], right[name], args.examples)
    payload = json.dumps(result, indent=2, ensure_ascii=False, default=str) + "\n"
    if args.output:
        with open(args.output, "w", encoding="utf-8") as handle:
            handle.write(payload)
    else:
        print(payload, end="")


if __name__ == "__main__":
    main()
