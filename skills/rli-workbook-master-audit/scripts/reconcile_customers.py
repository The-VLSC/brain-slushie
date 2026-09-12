#!/usr/bin/env python3
"""Create a review-only reconciliation package for two Customers Excel tables."""

from __future__ import annotations

import argparse
import json
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.table import Table, TableStyleInfo


def normalize_key(value: object) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Z0-9]+", " ", text.upper()).strip()


def serializable(value: object) -> object:
    if value.__class__.__name__ == "ArrayFormula":
        return {"formula_type": "ArrayFormula", "text": value.text, "ref": value.ref}
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def comparable(value: object) -> str:
    return json.dumps(serializable(value), ensure_ascii=False, sort_keys=True, default=str)


def read_customers(path: Path) -> tuple[list[str], list[dict[str, object]]]:
    workbook = load_workbook(path, read_only=False, data_only=False, keep_links=True)
    if "Clients" not in workbook.sheetnames:
        raise ValueError(f"{path}: missing Clients sheet")
    sheet = workbook["Clients"]
    if "Customers" not in sheet.tables:
        raise ValueError(f"{path}: missing Customers table")
    min_col, min_row, max_col, max_row = range_boundaries(sheet.tables["Customers"].ref)
    headers = [str(sheet.cell(min_row, col).value) for col in range(min_col, max_col + 1)]
    rows: list[dict[str, object]] = []
    for row_num in range(min_row + 1, max_row + 1):
        record = {
            header: serializable(sheet.cell(row_num, col).value)
            for header, col in zip(headers, range(min_col, max_col + 1))
        }
        if not any(value not in (None, "") for value in record.values()):
            continue
        record["_source_row"] = row_num
        record["_key"] = (
            normalize_key(record.get("CUSTOMER")),
            normalize_key(record.get("ADDRESS")),
        )
        rows.append(record)
    return headers, rows


def index_rows(rows: list[dict[str, object]]) -> dict[tuple[str, str], list[dict[str, object]]]:
    indexed: dict[tuple[str, str], list[dict[str, object]]] = defaultdict(list)
    for row in rows:
        indexed[row["_key"]].append(row)
    return indexed


def write_sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list[object]]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append([
            json.dumps(value, ensure_ascii=False, sort_keys=True)
            if isinstance(value, (dict, list))
            else value
            for value in row
        ])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(vertical="top")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, header in enumerate(headers, start=1):
        observed = [len(str(header))] + [len(str(sheet.cell(r, index).value or "")) for r in range(2, min(sheet.max_row, 200) + 1)]
        sheet.column_dimensions[get_column_letter(index)].width = min(max(observed) + 2, 55)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    if rows:
        table = Table(displayName=re.sub(r"[^A-Za-z0-9]", "", title)[:200], ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("master", type=Path)
    parser.add_argument("incoming", type=Path)
    parser.add_argument("--xlsx", type=Path, required=True)
    parser.add_argument("--json", type=Path, required=True)
    args = parser.parse_args()

    master_headers, master_rows = read_customers(args.master)
    incoming_headers, incoming_rows = read_customers(args.incoming)
    fields = list(dict.fromkeys(master_headers + incoming_headers))
    master_index = index_rows(master_rows)
    incoming_index = index_rows(incoming_rows)
    keys = sorted(set(master_index) | set(incoming_index))

    only_master: list[list[object]] = []
    only_incoming: list[list[object]] = []
    changed: list[list[object]] = []
    duplicates: list[list[object]] = []
    identical = 0
    field_counts: Counter[str] = Counter()

    for key in keys:
        left, right = master_index.get(key, []), incoming_index.get(key, [])
        display_name = (left or right)[0].get("CUSTOMER")
        display_address = (left or right)[0].get("ADDRESS")
        if len(left) > 1 or len(right) > 1:
            duplicates.append([display_name, display_address, len(left), len(right), ", ".join(str(x["_source_row"]) for x in left), ", ".join(str(x["_source_row"]) for x in right)])
            continue
        if not right:
            record = left[0]
            only_master.append([record.get("_source_row")] + [record.get(field) for field in fields])
            continue
        if not left:
            record = right[0]
            only_incoming.append([record.get("_source_row")] + [record.get(field) for field in fields])
            continue
        differences = 0
        for field in fields:
            master_value, incoming_value = left[0].get(field), right[0].get(field)
            if comparable(master_value) != comparable(incoming_value):
                differences += 1
                field_counts[field] += 1
                changed.append([display_name, display_address, field, master_value, incoming_value, left[0]["_source_row"], right[0]["_source_row"]])
        if differences == 0:
            identical += 1

    summary = {
        "master_rows": len(master_rows),
        "incoming_rows": len(incoming_rows),
        "identical_unique_keys": identical,
        "master_only_unique_keys": len(only_master),
        "incoming_only_unique_keys": len(only_incoming),
        "changed_field_records": len(changed),
        "ambiguous_duplicate_keys": len(duplicates),
        "field_difference_counts": dict(field_counts.most_common()),
    }
    args.json.write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")

    workbook = Workbook()
    workbook.remove(workbook.active)
    write_sheet(workbook, "Summary", ["Metric", "Value"], [[key, value] for key, value in summary.items() if key != "field_difference_counts"] + [["Field difference counts", json.dumps(summary["field_difference_counts"], ensure_ascii=False)]])
    write_sheet(workbook, "Only Master", ["Source Row"] + fields, only_master)
    write_sheet(workbook, "Only Incoming", ["Source Row"] + fields, only_incoming)
    write_sheet(workbook, "Changed Fields", ["Customer", "Address", "Field", "Master", "Incoming", "Master Row", "Incoming Row"], changed)
    write_sheet(workbook, "Duplicate Keys", ["Customer", "Address", "Master Count", "Incoming Count", "Master Rows", "Incoming Rows"], duplicates)
    workbook.save(args.xlsx)


if __name__ == "__main__":
    main()
